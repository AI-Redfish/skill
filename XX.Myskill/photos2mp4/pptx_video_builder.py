#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
pptx_video_builder.py - 批量图片 + 背景音乐 -> PPTX -> MP4 视频生成工具
=======================================================================

定位：作为 AI Skill 的底层脚本，由 AI 通过命令行调用，完成「批量图片生成视频」。

能力总览（对外暴露的命令）：
  build             图片组 + 背景音乐 -> PPTX（一图一页；每页切换效果；自动换片时间）
  set-transitions   修改已有 PPTX 的幻灯片切换方式 / 自动换片时间（支持逐页 JSON）
  list-transitions  输出切换效果目录；默认用本机 PowerPoint 实测哪些真正可用
  export            调用本机 PowerPoint 把 PPTX 导出为 MP4 视频（含背景音乐）
  all               build + export 一条龙
  selftest          离线自检（不需要 PowerPoint / pywin32），验证 XML 结构与几何计算

依赖：
  必装   pip install python-pptx
  导出   pip install pywin32        （仅 export / list-transitions 探测需要；需本机装有
                                     Microsoft PowerPoint 2010 及以上版本）
  可选   pip install mutagen       （仅 --auto-fit-music 读取音乐时长时需要）

典型工作流（AI 调用示例）：
  1) python pptx_video_builder.py list-transitions --json
  2) python pptx_video_builder.py build --images ./pics --music bgm.mp3 ^
         --output deck.pptx --transition fade --transition-duration 0.7 --advance 4
  3) python pptx_video_builder.py export --pptx deck.pptx --output video.mp4 --resolution 1080
  一条龙：
     python pptx_video_builder.py all --images ./pics --music bgm.mp3 ^
            --output video.mp4 --transition push --direction left

逐页切换 JSON（--transitions-json）格式（数组，元素字段均可选）：
  [
    {"slide": 1, "transition": "fade",    "duration": 0.8, "advance": 4.0},
    {"slide": 2, "transition": "push",    "direction": "left", "advance": 3.5},
    {"slide": 3, "transition": "wheel",   "options": {"spokes": "3"}}
  ]
  - "slide"      : 1 起始的页码；省略时按数组位置顺序对应（第 1 个未指定者对应第 1 页）
  - "transition" : 切换效果名（见 list-transitions）
  - "direction"  : 方向（left/right/up/down、horz/vert、in/out、throughblack 等，自动归一化）
  - "duration"   : 切换动画时长（秒）；精确时长需 PowerPoint 2010+（旧版自动降级为快/中/慢）
  - "advance"    : 本页停留秒数（自动换片时间；导出视频的节奏来源）
  - "options"    : 附加参数字典，如 {"spokes": "3"}、{"pattern": "hexagon"}

背景音乐说明：
  - 音乐以 OOXML 音频形状写入第 1 页：自动播放、跨全部幻灯片(numSld=999)、
    可循环、放映/导出视频时图标不可见（置于页面外）。
  - 推荐格式 mp3 / m4a / wav / wma。

退出码（便于 AI 判断结果）：0 成功；1 一般错误；2 参数/校验错误；
3 PowerPoint 不可用；4 导出超时/失败；130 用户中断。
所有人类可读日志走 stderr，JSON 数据走 stdout（--json 时）。
"""

import argparse
import glob as _glob
import json
import os
import re
import struct
import sys
import tempfile
import time
import zipfile
import zlib
from io import BytesIO
from xml.sax.saxutils import escape as _xml_escape

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as _RT
from pptx.opc.package import Part as _Part
from pptx.opc.packuri import PackURI as _PackURI
from pptx.oxml import parse_xml
from pptx.util import Emu

# 让中文日志在 Git Bash/管道环境下不乱码（交互式 cmd 控制台不受影响）
for _stream in (sys.stdout, sys.stderr):
    try:
        if _stream.encoding and _stream.encoding.lower().replace("-", "") != "utf8":
            _stream.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

# --------------------------------------------------------------------------
# 常量
# --------------------------------------------------------------------------

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"
P159_NS = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"

AUDIO_RT = getattr(_RT, "AUDIO",
                   "http://schemas.openxmlformats.org/officeDocument/2006/relationships/audio")
MEDIA_RT = "http://schemas.microsoft.com/office/2007/relationships/media"

AUDIO_CONTENT_TYPES = {
    ".mp3": "audio/mpeg",
    ".m4a": "audio/mp4",
    ".wav": "audio/x-wav",
    ".wma": "audio/x-ms-wma",
    ".aac": "audio/aac",
    ".flac": "audio/flac",
}

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}

EMU_PER_PX = 9525  # 96dpi

# 经典切换效果（ECMA-376，PowerPoint 2007+ 全部支持）
# tag: transition 子元素名；attr+values: direction 允许值；extra: 其它可设属性
CLASSIC_TRANSITIONS = {
    "blinds":    {"tag": "blinds",    "attr": "orient", "values": ["horz", "vert"]},
    "checker":   {"tag": "checker",   "attr": "orient", "values": ["horz", "vert"]},
    "circle":    {"tag": "circle"},
    "comb":      {"tag": "comb",      "attr": "orient", "values": ["horz", "vert"]},
    "cover":     {"tag": "cover",     "attr": "dir", "values": ["d", "l", "r", "u", "ld", "lu", "rd", "ru"]},
    "cut":       {"tag": "cut",       "attr": "thruBlk", "values": ["1"], "dir_alias": True},
    "diamond":   {"tag": "diamond"},
    "dissolve":  {"tag": "dissolve"},
    "fade":      {"tag": "fade",      "attr": "thruBlk", "values": ["1"], "dir_alias": True},
    "newsflash": {"tag": "newsflash"},
    "plus":      {"tag": "plus"},
    "pull":      {"tag": "pull",      "attr": "dir", "values": ["d", "l", "r", "u", "ld", "lu", "rd", "ru"]},
    "push":      {"tag": "push",      "attr": "dir", "values": ["d", "l", "r", "u"]},
    "random":    {"tag": "random"},
    "randombar": {"tag": "randomBar", "attr": "orient", "values": ["horz", "vert"]},
    "split":     {"tag": "split",     "attr": "dir", "values": ["in", "out"],
                  "extra": {"orient": ["horz", "vert"]}},
    "strips":    {"tag": "strips",    "attr": "dir", "values": ["ld", "lu", "rd", "ru"]},
    "wedge":     {"tag": "wedge"},
    "wheel":     {"tag": "wheel",     "attr": "spokes", "values": ["1", "2", "3", "4", "8"]},
    "wipe":      {"tag": "wipe",      "attr": "dir", "values": ["d", "l", "r", "u"]},
    "zoom":      {"tag": "zoom",      "attr": "dir", "values": ["in", "out"]},
}

# PowerPoint 2010+ 切换效果（写入时包裹在 mc:AlternateContent 中，旧版自动回退为 fade）
# required: schema 必填属性（缺失会导致 PowerPoint 拒绝打开文件！），自动补默认值
P14_TRANSITIONS = {
    "conveyor":   {"tag": "conveyor",   "attr": "dir", "values": ["l", "r"], "required": {"dir": "l"}},
    "doors":      {"tag": "doors"},
    "ferris":     {"tag": "ferris",     "attr": "dir", "values": ["l", "r"], "required": {"dir": "l"}},
    "flash":      {"tag": "flash"},
    "flip":       {"tag": "flip",       "attr": "dir", "values": ["l", "r"], "required": {"dir": "l"}},
    "flythrough": {"tag": "flythrough", "attr": "dir", "values": ["in", "out"],
                   "extra": {"hasBounce": ["0", "1"]}},
    "gallery":    {"tag": "gallery",    "attr": "dir", "values": ["l", "r"], "required": {"dir": "l"}},
    "glitter":    {"tag": "glitter",    "attr": "dir", "values": ["d", "l", "r", "u"],
                   "extra": {"pattern": ["diamond", "hexagon"]}},
    "honeycomb":  {"tag": "honeycomb"},
    "pan":        {"tag": "pan"},
    "prism":      {"tag": "prism"},
    "reveal":     {"tag": "reveal",     "attr": "thruBlk", "values": ["0", "1"], "dir_alias": True},
    "ripple":     {"tag": "ripple"},
    "shred":      {"tag": "shred",      "attr": "dir", "values": ["in", "out"],
                   "extra": {"pattern": ["strip", "rectangle"]}},
    "switch":     {"tag": "switch",     "attr": "dir", "values": ["l", "r"], "required": {"dir": "l"}},
    "vortex":     {"tag": "vortex",     "attr": "dir", "values": ["l", "r"]},
    "warp":       {"tag": "warp",       "attr": "dir", "values": ["in", "out"]},
    "window":     {"tag": "window"},
}

# PowerPoint 2019/365 的平滑切换（对纯图片幻灯片意义不大，列出仅供完整性）
MORPH_TRANSITIONS = {
    "morph": {"tag": "morph", "extra": {"option": ["byObject", "byWord", "byChar"]},
              "required": {"option": "byObject"}},
}

TIER_NS = {"classic": P_NS, "p14": P14_NS, "morph": P159_NS}
TIER_PREFIX = {"classic": "p", "p14": "p14", "morph": "p159"}
TIER_SINCE = {"classic": "PowerPoint 2007+", "p14": "PowerPoint 2010+",
              "morph": "PowerPoint 2019/365"}

DIR_ALIAS = {
    "left": "l", "right": "r", "up": "u", "down": "d", "top": "u", "bottom": "d",
    "horz": "horz", "horizontal": "horz", "vert": "vert", "vertical": "vert",
    "in": "in", "out": "out", "inward": "in", "outward": "out",
    "throughblack": "thruBlk", "thrublack": "thruBlk",
    "ld": "ld", "leftdown": "ld", "downleft": "ld",
    "lu": "lu", "leftup": "lu", "upleft": "lu",
    "rd": "rd", "rightdown": "rd", "downright": "rd",
    "ru": "ru", "rightup": "ru", "upright": "ru",
}

COLOR_NAMES = {
    "black": "000000", "white": "FFFFFF", "gray": "808080", "grey": "808080",
    "darkgray": "A9A9A9", "lightgray": "D3D3D3", "red": "FF0000", "green": "008000",
    "blue": "0000FF", "yellow": "FFFF00", "orange": "FFA500", "purple": "800080",
    "navy": "000080", "teal": "008080", "maroon": "800000", "silver": "C0C0C0",
    "olive": "808000", "lime": "00FF00", "aqua": "00FFFF", "cyan": "00FFFF",
    "fuchsia": "FF00FF", "magenta": "FF00FF", "brown": "A52A2A",
}

TAG_TRANSITION = "{%s}transition" % P_NS
TAG_TIMING = "{%s}timing" % P_NS
TAG_EXT_LST = "{%s}extLst" % P_NS
TAG_AC = "{%s}AlternateContent" % MC_NS


# --------------------------------------------------------------------------
# 基础工具
# --------------------------------------------------------------------------

class SkillError(Exception):
    """带退出码的业务错误。"""
    def __init__(self, message, exit_code=2):
        super().__init__(message)
        self.exit_code = exit_code


class PowerPointUnavailableError(SkillError):
    def __init__(self, message):
        super().__init__(message, exit_code=3)


def log(msg):
    sys.stderr.write(str(msg) + "\n")
    sys.stderr.flush()


def json_out(obj):
    print(json.dumps(obj, ensure_ascii=False, indent=2))


def natural_key(s):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", s)]


def _ensure_parent(path):
    parent = os.path.dirname(os.path.abspath(path))
    if parent:
        os.makedirs(parent, exist_ok=True)


def _localname(tag):
    return tag.rsplit("}", 1)[-1].lower() if isinstance(tag, str) else ""


def parse_color(value):
    v = str(value).strip().lower().lstrip("#")
    if v in COLOR_NAMES:
        return COLOR_NAMES[v]
    if re.fullmatch(r"[0-9a-f]{6}", v):
        return v.upper()
    if re.fullmatch(r"[0-9a-f]{3}", v):
        return "".join(c * 2 for c in v).upper()
    raise SkillError(f"无法识别的颜色: {value}（支持颜色名或 #RRGGBB）")


def parse_slide_size(value):
    """'16:9' | '4:3' | '16:10' | '1280x720'(像素,96dpi) -> (emu_w, emu_h)"""
    v = str(value).strip().lower()
    presets = {
        "16:9": (12192000, 6858000),
        "4:3": (9144000, 6858000),
        "16:10": (9144000, 5715000),
    }
    if v in presets:
        return presets[v]
    m = re.fullmatch(r"(\d+)[x*](\d+)", v)
    if m:
        w, h = int(m.group(1)), int(m.group(2))
        if w < 64 or h < 64 or w > 100000 or h > 100000:
            raise SkillError(f"幻灯片像素尺寸超出合理范围: {value}")
        return w * EMU_PER_PX, h * EMU_PER_PX
    raise SkillError(f"无法识别的幻灯片尺寸: {value}（支持 16:9 / 4:3 / 16:10 / 1280x720）")


def collect_images(args_list, recursive=False):
    """收集图片：支持目录、glob 通配、具体文件；自然排序、去重、过滤扩展名。"""
    if not args_list:
        raise SkillError("未提供图片（--images）")
    files = []
    for a in args_list:
        if os.path.isdir(a):
            if recursive:
                for root, _dirs, names in os.walk(a):
                    for nm in names:
                        if os.path.splitext(nm)[1].lower() in IMAGE_EXTS:
                            files.append(os.path.join(root, nm))
            else:
                for nm in os.listdir(a):
                    p = os.path.join(a, nm)
                    if os.path.isfile(p) and os.path.splitext(nm)[1].lower() in IMAGE_EXTS:
                        files.append(p)
        elif any(ch in a for ch in "*?["):
            matched = [p for p in _glob.glob(a) if os.path.isfile(p)]
            if not matched:
                log(f"[警告] 通配符未匹配到任何文件: {a}")
            files.extend(matched)
        elif os.path.isfile(a):
            files.append(a)
        else:
            raise SkillError(f"找不到图片文件: {a}")
    skipped = [f for f in files if os.path.splitext(f)[1].lower() not in IMAGE_EXTS]
    for f in skipped:
        log(f"[警告] 跳过不支持的图片格式: {f}（支持 {', '.join(sorted(IMAGE_EXTS))}）")
    files = [f for f in files if os.path.splitext(f)[1].lower() in IMAGE_EXTS]
    seen = set()
    unique = []
    for f in files:
        key = os.path.normcase(os.path.abspath(f))
        if key not in seen:
            seen.add(key)
            unique.append(f)
    unique.sort(key=lambda p: (natural_key(os.path.basename(p)), natural_key(p)))
    if not unique:
        raise SkillError("没有收集到任何可用图片")
    return unique


# --------------------------------------------------------------------------
# 切换效果：目录、归一化、XML 生成
# --------------------------------------------------------------------------

def _catalogs():
    return (("classic", CLASSIC_TRANSITIONS), ("p14", P14_TRANSITIONS),
            ("morph", MORPH_TRANSITIONS))


def lookup_transition(name):
    """按归一化名称查找，返回 (tier, key, entry)。找不到抛 SkillError。"""
    key = re.sub(r"[^a-z0-9]", "", str(name).strip().lower())
    if not key:
        raise SkillError("切换效果名为空")
    for tier, catalog in _catalogs():
        for k in catalog:
            if k == key:
                return tier, k, catalog[k]
    all_names = [k for _t, c in _catalogs() for k in c]
    raise SkillError(f"未知的切换效果: {name}（可用: {', '.join(all_names)}，"
                     f"或用 list-transitions 查看）")


def transition_attrs(tier, key, entry, direction, options):
    """把 direction/options 归一化并校验，返回属性 dict（写入子元素）。"""
    attrs = {}
    if direction is not None and str(direction).strip() != "":
        d = DIR_ALIAS.get(str(direction).strip().lower(), str(direction).strip().lower())
        if tier == "classic" and key == "split" and d in ("horz", "vert"):
            attrs["orient"] = d
        elif entry.get("dir_alias") and d == "thruBlk":
            attrs[entry["attr"]] = "1"
        elif entry.get("attr"):
            vals = entry["values"]
            if d not in vals:
                raise SkillError(f"切换效果 {key} 不接受方向 '{direction}'（允许: {', '.join(vals)}）")
            attrs[entry["attr"]] = d
        else:
            raise SkillError(f"切换效果 {key} 不支持方向参数")
    for k, v in (options or {}).items():
        allowed = dict(entry.get("extra") or {})
        if entry.get("attr") and k == entry["attr"]:
            allowed[k] = entry["values"]
        if k not in allowed:
            legal = ", ".join(sorted(allowed)) or "无"
            raise SkillError(f"切换效果 {key} 不支持选项 '{k}'（允许: {legal}）")
        if str(v) not in allowed[k]:
            raise SkillError(f"选项 {k}={v} 不合法（允许: {', '.join(allowed[k])}）")
        attrs[k] = str(v)
    # schema 必填属性：缺失会导致 PowerPoint 拒绝打开文件，自动补默认值
    for k, v in (entry.get("required") or {}).items():
        attrs.setdefault(k, v)
    return attrs


def _spd_for(duration_sec):
    if not duration_sec:
        return ""
    if duration_sec <= 0.75:
        return "fast"
    if duration_sec <= 1.5:
        return "med"
    return "slow"


def build_transition_element(spec, advance_ms=None):
    """根据切换描述生成可插入 p:sld 的元素；spec 含 transition/direction/duration/options。"""
    name = (spec.get("transition") or "none").strip()
    adv_attr = ' advClick="1"'
    if advance_ms:
        adv_attr = ' advClick="1" advTm="%d"' % int(advance_ms)

    norm = re.sub(r"[^a-z0-9]", "", str(name).lower())
    if norm in ("none", "null") or str(name).strip() == "无":
        return parse_xml('<p:transition xmlns:p="%s"%s/>' % (P_NS, adv_attr))

    tier, key, entry = lookup_transition(name)
    attrs = transition_attrs(tier, key, entry, spec.get("direction"), spec.get("options"))
    dur = spec.get("duration")
    child = _child_xml(tier, entry, attrs)

    if tier == "classic":
        if dur:
            spd = _spd_for(dur)
            spd_a = (' spd="%s"' % spd) if spd else ""
            choice = ('<p:transition xmlns:p="%s"%s p14:dur="%d"%s>%s</p:transition>'
                      % (P_NS, spd_a, int(round(dur * 1000)), adv_attr, child))
            fallback = ('<p:transition xmlns:p="%s"%s%s>%s</p:transition>'
                        % (P_NS, spd_a, adv_attr, child))
            xml = _ac_wrap("p14", choice, fallback)
        else:
            xml = '<p:transition xmlns:p="%s"%s>%s</p:transition>' % (P_NS, adv_attr, child)
        return parse_xml(xml)

    # p14 / morph：必须包裹 mc:AlternateContent，旧版回退 fade
    req = "p14" if tier == "p14" else "p159"
    spd = _spd_for(dur)
    spd_a = (' spd="%s"' % spd) if spd else ""
    pdur = (' p14:dur="%d"' % int(round(dur * 1000))) if dur else ""
    choice = '<p:transition xmlns:p="%s"%s%s%s>%s</p:transition>' % (
        P_NS, spd_a, pdur, adv_attr, child)
    fallback = '<p:transition xmlns:p="%s"%s%s><p:fade/></p:transition>' % (
        P_NS, spd_a, adv_attr)
    return parse_xml(_ac_wrap(req, choice, fallback))


def _child_xml(tier, entry, attrs):
    prefix = TIER_PREFIX[tier]
    attr_s = "".join(' %s="%s"' % (k, v) for k, v in sorted(attrs.items()))
    return "<%s:%s%s/>" % (prefix, entry["tag"], attr_s)


def _ac_wrap(requires, choice_inner, fallback_inner):
    return ('<mc:AlternateContent xmlns:mc="%s" xmlns:p14="%s" xmlns:p159="%s">'
            '<mc:Choice Requires="%s">%s</mc:Choice>'
            '<mc:Fallback>%s</mc:Fallback></mc:AlternateContent>'
            % (MC_NS, P14_NS, P159_NS, requires, choice_inner, fallback_inner))


def _remove_existing_transition(sld):
    for child in list(sld):
        if child.tag == TAG_TRANSITION:
            sld.remove(child)
        elif child.tag == TAG_AC and child.find(".//" + TAG_TRANSITION) is not None:
            sld.remove(child)


def _existing_advTm(sld):
    """读取现有切换上的 advTm（毫秒），无则 None。"""
    tr = None
    for child in sld:
        if child.tag == TAG_TRANSITION:
            tr = child
            break
        if child.tag == TAG_AC:
            found = child.find(".//" + TAG_TRANSITION)
            if found is not None:
                tr = found
                break
    if tr is None:
        return None
    v = tr.get("advTm")
    try:
        return int(v) if v is not None else None
    except ValueError:
        return None


def set_slide_transition(slide, spec, advance_ms="keep"):
    """设置单页切换。advance_ms: 毫秒数 / None(清除自动换片) / 'keep'(沿用已有值)。"""
    sld = slide._element
    if advance_ms == "keep":
        advance_ms = _existing_advTm(sld)  # 必须在删除前读取旧值
    _remove_existing_transition(sld)
    el = build_transition_element(spec, advance_ms)
    anchor = next((c for c in sld if c.tag in (TAG_TIMING, TAG_EXT_LST)), None)
    if anchor is not None:
        anchor.addprevious(el)
    else:
        sld.append(el)


def _entry_params(entry):
    """生成切换效果的人类可读参数说明。"""
    params = []
    if entry.get("attr"):
        if entry.get("dir_alias"):
            params.append("direction: throughblack")
        else:
            params.append("direction: %s" % "|".join(entry["values"]))
    for k, vals in (entry.get("extra") or {}).items():
        params.append("%s: %s" % (k, "|".join(vals)))
    return "; ".join(params) if params else "-"


def catalog_rows():
    rows = []
    for tier, catalog in _catalogs():
        for key, entry in catalog.items():
            rows.append({"name": key, "tier": tier, "since": TIER_SINCE[tier],
                         "params": _entry_params(entry)})
    return rows


# --------------------------------------------------------------------------
# 背景音乐（OOXML 音频形状 + 跨页 timing）
# --------------------------------------------------------------------------

def make_png_bytes(width, height, rgb):
    """纯 stdlib 生成合法 PNG（truecolor 8bit）。"""
    def chunk(tag, data):
        return (struct.pack(">I", len(data)) + tag + data
                + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF))
    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def make_wav_bytes(seconds=1.0, rate=8000):
    """纯 stdlib 生成静音 WAV（自检用）。"""
    import wave as _wave
    buf = BytesIO()
    w = _wave.open(buf, "wb")
    w.setnchannels(1)
    w.setsampwidth(2)
    w.setframerate(rate)
    w.writeframes(b"\x00\x00" * int(rate * seconds))
    w.close()
    return buf.getvalue()


def music_duration(path):
    try:
        import mutagen
    except ImportError:
        raise SkillError("--auto-fit-music 需要 mutagen：pip install mutagen")
    m = mutagen.File(path)
    if m is None or getattr(m, "info", None) is None:
        raise SkillError(f"无法读取音乐时长: {path}")
    length = float(m.info.length)
    if length <= 0:
        raise SkillError(f"音乐时长异常: {path}")
    return length


def _new_blob_part(package, partname, content_type, blob):
    try:
        return _Part(_PackURI(partname), content_type, package=package, blob=blob)
    except TypeError:  # 兼容旧版 python-pptx 的位置参数签名
        return _Part(_PackURI(partname), content_type, package, blob)


def _next_shape_id(sld):
    max_id = 1
    for el in sld.iter():
        v = el.get("id")
        if v and v.isdigit():
            max_id = max(max_id, int(v))
    return max_id + 1


def add_background_music(slide, music_path, loop=True, volume=100):
    """把音乐作为背景音乐写入 slide：自动播放、跨全部页、可循环、图标在页面外不可见。"""
    if not os.path.isfile(music_path):
        raise SkillError(f"找不到音乐文件: {music_path}")
    ext = os.path.splitext(music_path)[1].lower()
    if ext not in AUDIO_CONTENT_TYPES:
        raise SkillError(f"不支持的音乐格式: {ext}（支持 {', '.join(sorted(AUDIO_CONTENT_TYPES))}）")
    if not 1 <= int(volume) <= 100:
        raise SkillError("--music-volume 取值范围 1-100")

    with open(music_path, "rb") as f:
        blob = f.read()
    package = slide.part.package

    existing = {str(p.partname) for p in package.iter_parts()}
    idx = 1
    while "/ppt/media/media%d%s" % (idx, ext) in existing:
        idx += 1
    partname = "/ppt/media/media%d%s" % (idx, ext)
    audio_part = _new_blob_part(package, partname, AUDIO_CONTENT_TYPES[ext], blob)
    rid_audio = slide.part.relate_to(audio_part, AUDIO_RT)
    rid_media = slide.part.relate_to(audio_part, MEDIA_RT)

    icon_rid = slide.part.get_or_add_image_part(BytesIO(make_png_bytes(1, 1, (89, 89, 89))))[1]

    sld = slide._element
    spid = _next_shape_id(sld)
    sp_tree = sld.find("{%s}cSld/{%s}spTree" % (P_NS, P_NS))
    if sp_tree is None:
        raise SkillError("幻灯片 XML 结构异常（缺少 spTree）")

    pic_xml = (
        '<p:pic xmlns:p="%s" xmlns:a="%s" xmlns:r="%s" xmlns:p14="%s">'
        '<p:nvPicPr>'
        '<p:cNvPr id="%d" name="%s">'
        '<a:hlinkClick r:id="" action="ppaction://media"/>'
        '</p:cNvPr>'
        '<p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr>'
        '<p:nvPr>'
        '<a:audioFile r:link="%s"/>'
        '<p:extLst><p:ext uri="{DAA4B4D4-6D71-4841-9C94-3DE7FCFB9230}">'
        '<p14:media r:embed="%s"/>'
        '</p:ext></p:extLst>'
        '</p:nvPr>'
        '</p:nvPicPr>'
        '<p:blipFill><a:blip r:embed="%s"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
        '<p:spPr><a:xfrm><a:off x="-10000" y="-10000"/><a:ext cx="1000" cy="1000"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
        '</p:pic>'
    ) % (P_NS, A_NS, R_NS, P14_NS, spid,
         _xml_escape("Background Music (%s)" % os.path.basename(music_path),
                     {'"': "&quot;"}),
         rid_audio, rid_media, icon_rid)
    sp_tree.append(parse_xml(pic_xml))

    repeat_attr = ' repeatCount="indefinite"' if loop else ""
    timing_xml = (
        '<p:timing xmlns:p="%s" xmlns:a="%s">'
        '<p:tnLst><p:par>'
        '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
        '<p:childTnLst>'
        '<p:audio>'
        '<p:cMediaNode vol="%d" numSld="999" showWhenStopped="0">'
        '<p:cTn id="2"%s fill="hold" display="0" masterRel="sameClick">'
        '<p:stCondLst><p:cond delay="0"/></p:stCondLst>'
        '<p:endCondLst><p:cond evt="onStopAudio" delay="0">'
        '<p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:endCondLst>'
        '</p:cTn>'
        '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl>'
        '</p:cMediaNode>'
        '</p:audio>'
        '</p:childTnLst>'
        '</p:cTn>'
        '</p:par></p:tnLst>'
        '</p:timing>'
    ) % (P_NS, A_NS, int(volume) * 1000, repeat_attr, spid)
    set_slide_timing(slide, timing_xml)


def set_slide_timing(slide, timing_xml):
    sld = slide._element
    for child in list(sld):
        if child.tag == TAG_TIMING:
            sld.remove(child)
    el = parse_xml(timing_xml)
    anchor = next((c for c in sld if c.tag == TAG_EXT_LST), None)
    if anchor is not None:
        anchor.addprevious(el)
    else:
        sld.append(el)


# --------------------------------------------------------------------------
# 构建 PPTX
# --------------------------------------------------------------------------

def _set_background(slide, hex6, sw, sh):
    try:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(hex6)
    except Exception:
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), Emu(sw), Emu(sh))
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(hex6)
        shp.line.fill.background()
        try:
            shp.shadow.inherit = False
        except Exception:
            pass


def _place_picture(slide, image_path, fit, sw, sh):
    pic = slide.shapes.add_picture(image_path, Emu(0), Emu(0))
    iw, ih = int(pic.width), int(pic.height)
    if iw <= 0 or ih <= 0:
        raise SkillError(f"图片尺寸异常: {image_path}")
    if fit == "stretch":
        w, h = sw, sh
    elif fit == "cover":
        s = max(sw / iw, sh / ih)
        w, h = max(1, int(round(iw * s))), max(1, int(round(ih * s)))
    else:  # contain
        s = min(sw / iw, sh / ih)
        w, h = max(1, int(round(iw * s))), max(1, int(round(ih * s)))
    pic.width, pic.height = Emu(w), Emu(h)
    pic.left, pic.top = Emu((sw - w) // 2), Emu((sh - h) // 2)


def parse_transitions_json(path):
    try:
        with open(path, "r", encoding="utf-8-sig") as f:
            data = json.load(f)
    except OSError as e:
        raise SkillError(f"无法读取 --transitions-json 文件: {path} ({e})")
    except ValueError as e:
        raise SkillError(f"--transitions-json 不是合法 JSON: {e}")
    if isinstance(data, dict):
        data = [data]
    if not isinstance(data, list):
        raise SkillError("--transitions-json 必须是 JSON 数组（或单个对象）")
    allowed = {"slide", "transition", "direction", "duration", "advance", "options"}
    for i, item in enumerate(data):
        if not isinstance(item, dict):
            raise SkillError(f"--transitions-json 第 {i + 1} 项必须是对象")
        unknown = set(item) - allowed
        if unknown:
            raise SkillError(f"--transitions-json 第 {i + 1} 项含未知字段: {', '.join(sorted(unknown))}"
                             f"（允许: {', '.join(sorted(allowed))}）")
    return data


def build_pptx(images, output, music=None, transition="fade", direction=None,
               duration=None, advance=5.0, options=None, per_slide=None,
               fit="contain", background="black", slide_size="16:9",
               music_loop=True, music_volume=100, auto_fit_music=False):
    """核心构建：图片组(+音乐) -> PPTX。返回信息 dict。"""
    if fit not in ("contain", "cover", "stretch"):
        raise SkillError(f"未知的图片适配方式: {fit}")
    if not images:
        raise SkillError("没有可用的图片")
    if music and auto_fit_music:
        total = music_duration(music)
        sec = total / len(images)
        log(f"[信息] 音乐时长 {total:.1f}s，自动分配每页 {sec:.2f}s")
    sw, sh = parse_slide_size(slide_size)
    bg = parse_color(background)
    per_slide = per_slide or []

    n = len(images)
    specs = [{"transition": transition, "direction": direction, "duration": duration,
              "advance": advance, "options": options} for _ in range(n)]
    auto = 0
    for item in per_slide:
        if item.get("slide") is not None:
            idx = int(item["slide"]) - 1
        else:
            auto += 1
            idx = auto - 1
        if not 0 <= idx < n:
            raise SkillError(f"--transitions-json 指定的页码超出范围: {idx + 1}（共 {n} 页）")
        for k in ("transition", "direction", "duration", "advance", "options"):
            if item.get(k) is not None:
                specs[idx][k] = item[k]
    if music and auto_fit_music:
        sec = music_duration(music) / n
        for s in specs:
            s["advance"] = sec

    prs = Presentation()
    prs.slide_width = Emu(sw)
    prs.slide_height = Emu(sh)
    layout = prs.slide_layouts[6]

    for i, img in enumerate(images):
        slide = prs.slides.add_slide(layout)
        _set_background(slide, bg, sw, sh)
        _place_picture(slide, img, fit, sw, sh)
        set_slide_transition(slide, specs[i],
                             int(round(float(specs[i]["advance"] or 0) * 1000)))

    if music:
        add_background_music(prs.slides[0], music, loop=music_loop,
                             volume=music_volume)

    _ensure_parent(output)
    prs.save(output)
    return {"output": output, "slides": n, "music": music,
            "size_bytes": os.path.getsize(output)}


# --------------------------------------------------------------------------
# PowerPoint COM：版本探测 / 切换效果实测 / MP4 导出
# --------------------------------------------------------------------------

def _require_pywin32():
    try:
        import win32com.client  # noqa: F401
    except ImportError:
        raise PowerPointUnavailableError("需要 pywin32：pip install pywin32")


def _new_powerpoint():
    """创建独立 PowerPoint 实例（DispatchEx），不影响用户已打开的窗口。"""
    _require_pywin32()
    import win32com.client
    try:
        app = win32com.client.DispatchEx("PowerPoint.Application")
    except Exception as e:
        raise PowerPointUnavailableError(
            f"无法启动 PowerPoint COM（{e}）。请确认本机已安装 Microsoft PowerPoint 2010+。")
    return app


def _quit_app(app):
    try:
        app.Quit()
    except Exception:
        pass


def powerpoint_version():
    app = _new_powerpoint()
    try:
        return str(app.Version)
    finally:
        _quit_app(app)


def _effective_transition_localname(slide):
    """读取幻灯片实际生效的切换子元素 localname（小写），无切换返回 None。"""
    sld = slide._element
    for child in sld:
        if child.tag == TAG_AC:
            for c in child:
                if c.tag == "{%s}Choice" % MC_NS:
                    tr = c.find(".//" + TAG_TRANSITION)
                    if tr is not None:
                        kids = list(tr)
                        return _localname(kids[0].tag) if kids else None
    for child in sld:
        if child.tag == TAG_TRANSITION:
            kids = list(child)
            return _localname(kids[0].tag) if kids else None
    return None


def verify_transitions(workdir=None):
    """写探针 PPTX（每个效果一页）-> 让本机 PowerPoint 重新保存 -> 比对 XML 存活情况。

    整体探针打不开时（个别效果触发 PowerPoint 的硬校验失败），自动降级为
    逐个效果单独探测，确保任何版本 PowerPoint 都能得出完整支持矩阵。
    """
    workdir = workdir or tempfile.mkdtemp(prefix="pptx_probe_")
    os.makedirs(workdir, exist_ok=True)
    entries = [(tier, key) for tier, catalog in _catalogs() for key in catalog]
    probe_png = os.path.join(workdir, "probe.png")
    with open(probe_png, "wb") as f:
        f.write(make_png_bytes(64, 64, (200, 60, 60)))

    def _expected(tier, key):
        return "morph" if tier == "morph" else _catalog_entry(tier, key)["tag"].lower()

    def _result(tier, key, observed, note=None):
        r = {"name": key, "tier": tier, "since": TIER_SINCE[tier],
             "params": _entry_params(_catalog_entry(tier, key)),
             "expected": _expected(tier, key), "observed": observed,
             "supported": observed == _expected(tier, key)}
        if note:
            r["note"] = note
        return r

    app = _new_powerpoint()
    try:
        try:
            app.DisplayAlerts = 1  # ppAlertsNone
        except Exception:
            pass
        version = str(app.Version)

        def _open_save_close(src, dst):
            p = None
            try:
                p = app.Presentations.Open(src, False, False, False)
                p.SaveAs(dst)
            finally:
                try:
                    p.Close()
                except Exception:
                    pass

        results = []
        probe_pptx = os.path.join(workdir, "probe.pptx")
        saved_pptx = os.path.join(workdir, "probe_saved.pptx")
        build_pptx([probe_png] * len(entries), probe_pptx,
                   transition=entries[0][1], advance=1.0,
                   per_slide=[{"slide": i + 1, "transition": key, "advance": 1.0}
                              for i, (_tier, key) in enumerate(entries)])
        try:
            _open_save_close(os.path.abspath(probe_pptx), os.path.abspath(saved_pptx))
            prs = Presentation(saved_pptx)
            for i, slide in enumerate(prs.slides):
                tier, key = entries[i]
                results.append(_result(tier, key,
                                       _effective_transition_localname(slide)))
        except Exception:
            results = []

        if len(results) != len(entries):
            # 降级：逐个效果单独探测
            results = []
            for i, (tier, key) in enumerate(entries):
                single = os.path.abspath(os.path.join(workdir, "single_%02d.pptx" % i))
                single_saved = os.path.abspath(
                    os.path.join(workdir, "single_saved_%02d.pptx" % i))
                observed = None
                note = None
                try:
                    build_pptx([probe_png], single, transition=key, advance=1.0)
                    _open_save_close(single, single_saved)
                    prs2 = Presentation(single_saved)
                    observed = _effective_transition_localname(prs2.slides[0])
                except Exception:
                    observed = None
                    note = "open_failed"
                results.append(_result(tier, key, observed, note))
        return {"powerpoint_version": version, "results": results}
    finally:
        _quit_app(app)


def _catalog_entry(tier, key):
    for t, catalog in _catalogs():
        if t == tier and key in catalog:
            return catalog[key]
    raise KeyError(key)


def _deck_has_timings(pptx_path):
    """粗检：任意幻灯片带 advTm 自动换片时间则返回 True。"""
    try:
        with zipfile.ZipFile(pptx_path) as z:
            for name in z.namelist():
                if re.match(r"ppt/slides/slide\d+\.xml$", name):
                    if b"advTm=" in z.read(name):
                        return True
    except Exception:
        return False
    return False


def export_mp4(pptx_path, mp4_path, resolution=720, fps=30, quality=85,
               timeout=1800, slide_seconds=None):
    """调用本机 PowerPoint 将 PPTX 导出为 MP4。返回信息 dict。"""
    pptx_path = os.path.abspath(pptx_path)
    mp4_path = os.path.abspath(mp4_path)
    if not os.path.isfile(pptx_path):
        raise SkillError(f"找不到 PPTX 文件: {pptx_path}")
    if not pptx_path.lower().endswith(".pptx"):
        raise SkillError(f"输入必须是 .pptx: {pptx_path}")
    if not mp4_path.lower().endswith(".mp4"):
        raise SkillError(f"输出必须是 .mp4: {mp4_path}")
    if not 1 <= int(quality) <= 100:
        raise SkillError("--quality 取值范围 1-100")

    if not _deck_has_timings(pptx_path) and not slide_seconds:
        log("[警告] 所有幻灯片都没有自动换片时间(advTm)，导出的视频可能停在第 1 页。"
            "建议 build 时加 --advance 秒数，或 export 时加 --slide-seconds。")

    _ensure_parent(mp4_path)
    if os.path.exists(mp4_path):
        try:
            os.remove(mp4_path)
        except OSError as e:
            raise SkillError(f"输出文件已存在且无法删除（可能正被播放器占用）: {mp4_path} ({e})")

    app = None
    pres = None
    try:
        app = _new_powerpoint()
        try:
            app.DisplayAlerts = 1  # ppAlertsNone，避免弹窗阻塞
        except Exception:
            pass
        log(f"[信息] 打开演示文稿: {pptx_path}")
        try:
            pres = app.Presentations.Open(pptx_path, False, False, False)
            sec = int(slide_seconds) if slide_seconds else -1
            pres.CreateVideo(mp4_path, True, sec, int(resolution), int(fps), int(quality))
        except Exception as e:
            raise SkillError(
                f"PowerPoint 导出操作失败（{e}）。请确认已安装 PowerPoint 2010+，"
                "且文件未被其他程序占用。", 3)
        log(f"[信息] 开始导出视频（{resolution}p / {fps}fps / 质量 {quality}）: {mp4_path}")

        t0 = time.time()
        last_size = -1
        stable = 0
        last_log = 0.0
        while True:
            time.sleep(2)
            elapsed = time.time() - t0
            try:
                status = int(pres.CreateVideoStatus)
            except Exception:
                status = -1
            size = os.path.getsize(mp4_path) if os.path.exists(mp4_path) else 0
            if status == 3:  # ppMediaTaskStatusDone
                break
            if status == 4:  # ppMediaTaskStatusFailed
                raise SkillError("PowerPoint 报告视频导出失败（status=Failed），"
                                 "请检查幻灯片中的媒体是否可用。", 4)
            if timeout and elapsed > timeout:
                raise SkillError(f"导出超时（>{int(timeout)}s，status={status}）。"
                                 "可加大 --timeout 后重试。", 4)
            if status == 0 and elapsed > 20 and size > 0:
                stable = stable + 1 if size == last_size else 0
                if stable >= 3:
                    log("[警告] 状态未报告完成，但输出文件大小已稳定，视为完成。")
                    break
            elif size != last_size:
                stable = 0
            if size != last_size or elapsed - last_log > 30:
                last_size = size
                last_log = elapsed
                log(f"[信息]   导出中... {elapsed:.0f}s, status={status}, "
                    f"已写入 {size / 1048576:.1f} MB")
        size = os.path.getsize(mp4_path) if os.path.exists(mp4_path) else 0
        if size <= 0:
            raise SkillError("导出结束但输出文件为空。", 1)
        log(f"[信息] 导出完成: {mp4_path} ({size / 1048576:.1f} MB, "
            f"用时 {time.time() - t0:.0f}s)")
        return {"output": mp4_path, "bytes": size,
                "seconds": round(time.time() - t0, 1),
                "powerpoint_version": str(app.Version)}
    finally:
        if pres is not None:
            try:
                pres.Close()
            except Exception:
                pass
        if app is not None:
            _quit_app(app)


# --------------------------------------------------------------------------
# CLI 命令
# --------------------------------------------------------------------------

def _parse_opts(opt_list):
    """把 ['spokes=3', 'pattern=hexagon'] 解析为 dict。"""
    opts = {}
    for item in (opt_list or []):
        if "=" not in item:
            raise SkillError(f"--opt 需要 键=值 形式: {item}")
        k, v = item.split("=", 1)
        opts[k.strip()] = v.strip()
    return opts


def cmd_build(args):
    images = collect_images(args.images, recursive=args.recursive)
    per_slide = parse_transitions_json(args.transitions_json) if args.transitions_json else None
    if args.advance < 0:
        raise SkillError("--advance 不能为负数")
    info = build_pptx(images, args.output, music=args.music,
                      transition=args.transition, direction=args.direction,
                      duration=args.transition_duration, advance=args.advance,
                      options=_parse_opts(args.opt), per_slide=per_slide,
                      fit=args.fit, background=args.background,
                      slide_size=args.slide_size, music_loop=not args.music_no_loop,
                      music_volume=args.music_volume,
                      auto_fit_music=args.auto_fit_music)
    log(f"[OK] 已生成 PPTX: {info['output']}（{info['slides']} 页, "
        f"{info['size_bytes'] / 1048576:.2f} MB, 音乐: {info['music'] or '无'}）")
    if args.json:
        json_out(info)
    return 0


def cmd_set_transitions(args):
    if not args.transition and not args.transitions_json:
        raise SkillError("需要提供 --transition 或 --transitions-json 之一")
    per_slide = parse_transitions_json(args.transitions_json) if args.transitions_json else None
    if not os.path.isfile(args.pptx):
        raise SkillError(f"找不到 PPTX 文件: {args.pptx}")
    prs = Presentation(args.pptx)
    n = len(prs.slides._sldIdLst)
    specs = [{"transition": args.transition, "direction": args.direction,
              "duration": args.transition_duration, "advance": None,
              "options": _parse_opts(args.opt)} for _ in range(n)]
    if per_slide:
        auto = 0
        for item in per_slide:
            if item.get("slide") is not None:
                idx = int(item["slide"]) - 1
            else:
                auto += 1
                idx = auto - 1
            if not 0 <= idx < n:
                raise SkillError(f"--transitions-json 页码超出范围: {idx + 1}（共 {n} 页）")
            for k in ("transition", "direction", "duration", "advance", "options"):
                if item.get(k) is not None:
                    specs[idx][k] = item[k]
    advance_ms = "keep"
    if args.advance is not None:
        advance_ms = int(round(float(args.advance) * 1000))
    count = 0
    for i, slide in enumerate(prs.slides):
        spec = specs[i]
        if not spec.get("transition") and spec.get("advance") is not None:
            spec["transition"] = "none"  # 只改换片时间也需要写到 transition 元素上
        if spec.get("transition"):
            ms = advance_ms
            if spec.get("advance") is not None:
                ms = int(round(float(spec["advance"]) * 1000))
            set_slide_transition(slide, spec, ms)
            count += 1
    prs.save(args.pptx)
    log(f"[OK] 已更新 {count}/{n} 页的切换方式: {args.pptx}")
    if args.json:
        json_out({"pptx": args.pptx, "slides_updated": count, "slides_total": n})
    return 0


def cmd_list_transitions(args):
    rows = catalog_rows()
    if args.no_verify:
        if args.json:
            json_out({"powerpoint_version": None, "verified": False, "results": rows})
        else:
            _print_catalog_table(rows)
            log("[提示] 已跳过本机 PowerPoint 实测（--no-verify）。")
        return 0
    try:
        info = verify_transitions()
    except PowerPointUnavailableError as e:
        log(f"[警告] 无法连接本机 PowerPoint：{e}")
        if args.json:
            json_out({"powerpoint_version": None, "verified": False,
                      "error": str(e), "results": rows})
        else:
            _print_catalog_table(rows)
            log("[提示] 以上为理论支持列表（未经本机实测）。")
        return 3
    if args.json:
        json_out({"powerpoint_version": info["powerpoint_version"],
                  "verified": True, "results": info["results"]})
    else:
        _print_catalog_table(info["results"], verified=True,
                             version=info["powerpoint_version"])
    return 0


def _print_catalog_table(rows, verified=False, version=None):
    if verified:
        log(f"[信息] 本机 PowerPoint 版本: {version}，以下为实测结果：")
    name_w = max(len(r["name"]) for r in rows) + 2
    tier_w = max(len(r["tier"]) for r in rows) + 2
    header = "%-*s %-*s %-18s %-28s %s" % (
        name_w, "名称", tier_w, "级别", "要求版本", "参数", "本机实测")
    log(header)
    log("-" * len(header))
    for r in rows:
        support = "-"
        if verified:
            support = "[支持]" if r.get("supported") else "[不支持]"
        param_s = r["params"] if len(r["params"]) <= 28 else r["params"][:27] + "~"
        log("%-*s %-*s %-18s %-28s %s" % (
            name_w, r["name"], tier_w, r["tier"], r["since"], param_s, support))


def cmd_export(args):
    info = export_mp4(args.pptx, args.output, resolution=args.resolution,
                      fps=args.fps, quality=args.quality, timeout=args.timeout,
                      slide_seconds=args.slide_seconds)
    if args.json:
        json_out(info)
    return 0


def cmd_all(args):
    images = collect_images(args.images, recursive=args.recursive)
    per_slide = parse_transitions_json(args.transitions_json) if args.transitions_json else None
    output_mp4 = os.path.abspath(args.output)
    pptx_path = os.path.abspath(args.pptx) if args.pptx else \
        os.path.splitext(output_mp4)[0] + ".pptx"
    info = build_pptx(images, pptx_path, music=args.music,
                      transition=args.transition, direction=args.direction,
                      duration=args.transition_duration, advance=args.advance,
                      options=_parse_opts(args.opt), per_slide=per_slide,
                      fit=args.fit, background=args.background,
                      slide_size=args.slide_size, music_loop=not args.music_no_loop,
                      music_volume=args.music_volume,
                      auto_fit_music=args.auto_fit_music)
    log(f"[OK] 中间产物 PPTX: {info['output']}（{info['slides']} 页）")
    video = export_mp4(pptx_path, output_mp4, resolution=args.resolution,
                       fps=args.fps, quality=args.quality, timeout=args.timeout,
                       slide_seconds=args.slide_seconds)
    result = {"pptx": info["output"], "video": video["output"],
              "bytes": video["bytes"], "seconds": video["seconds"],
              "slides": info["slides"]}
    log(f"[OK] 完成：{result['pptx']} -> {result['video']}")
    if args.json:
        json_out(result)
    return 0


# --------------------------------------------------------------------------
# selftest：离线自检（不需要 PowerPoint / pywin32）
# --------------------------------------------------------------------------

def _selftest_make_assets(tmp):
    colors = [(200, 60, 60), (60, 200, 60), (60, 60, 200), (200, 200, 60)]
    paths = []
    for i, c in enumerate(colors):
        p = os.path.join(tmp, "img%d.png" % (i + 1))
        with open(p, "wb") as f:
            f.write(make_png_bytes(64, 64, c))
        paths.append(p)
    wav = os.path.join(tmp, "bgm.wav")
    with open(wav, "wb") as f:
        f.write(make_wav_bytes(2.0))
    square = os.path.join(tmp, "square.png")
    with open(square, "wb") as f:
        f.write(make_png_bytes(1000, 1000, (10, 10, 10)))
    return paths, wav, square


def _slide_transition_localname(slide):
    return _effective_transition_localname(slide)


def _slide_advTm(slide):
    from pptx.oxml.ns import qn
    sld = slide._element
    for child in sld:
        if child.tag == qn("p:transition"):
            return child.get("advTm")
        if child.tag == TAG_AC:
            tr = child.find(".//" + TAG_TRANSITION)
            if tr is not None:
                return tr.get("advTm")
    return None


def _run_tests(tmp):
    failures = []

    def check(name, fn):
        try:
            fn()
            log(f"  [OK] {name}")
        except Exception as e:  # noqa: BLE001
            import traceback
            failures.append((name, e))
            log(f"  [FAIL] {name}: {e!r}")
            traceback.print_exc(file=sys.stderr)

    images, wav, square = _selftest_make_assets(tmp)

    # ---- T1 基础构建：音乐 + 切换 + 自动换片 -------------------------------
    def t1():
        out = os.path.join(tmp, "t1.pptx")
        build_pptx(images, out, music=wav, transition="fade", duration=0.7,
                   advance=2.0, fit="contain", background="#101010")
        assert os.path.isfile(out), "输出文件不存在"
        with zipfile.ZipFile(out) as z:
            names = z.namelist()
            assert any(n.startswith("ppt/media/media") and n.endswith(".wav")
                       for n in names), "缺少音频 part"
            ct = z.read("[Content_Types].xml").decode("utf-8")
            assert "audio/x-wav" in ct, "[Content_Types].xml 缺少音频类型"
            s1 = z.read("ppt/slides/slide1.xml").decode("utf-8")
            assert "audioFile" in s1 and "numSld=\"999\"" in s1, "第 1 页缺少背景音乐结构"
            assert 'spTgt spid=' in s1, "timing 缺少音频目标"
            rels = z.read("ppt/slides/_rels/slide1.xml.rels").decode("utf-8")
            assert "relationships/audio" in rels, "缺少 audio 关系"
            assert "office/2007/relationships/media" in rels, "缺少 media 关系"
            for i in (2, 3, 4):
                sx = z.read("ppt/slides/slide%d.xml" % i).decode("utf-8")
                assert "audioFile" not in sx, f"第 {i} 页不应有音频"
        prs = Presentation(out)  # python-pptx 能重新打开
        assert len(prs.slides._sldIdLst) == 4
        assert _slide_transition_localname(prs.slides[0]) == "fade"
        assert _slide_advTm(prs.slides[0]) == "2000", "advTm 不正确"
    check("T1 基础构建（音乐/切换/换片时间/包结构）", t1)

    # ---- T2 全部经典切换效果可生成且结构正确 --------------------------------
    def t2():
        keys = list(CLASSIC_TRANSITIONS)
        out = os.path.join(tmp, "t2.pptx")
        per = [{"slide": i + 1, "transition": k, "advance": 1.0}
               for i, k in enumerate(keys)]
        build_pptx([images[0]] * len(keys), out, transition=keys[0],
                   advance=1.0, per_slide=per)
        prs = Presentation(out)
        for i, k in enumerate(keys):
            ln = _slide_transition_localname(prs.slides[i])
            assert ln == CLASSIC_TRANSITIONS[k]["tag"].lower(), \
                f"{k}: 期望 {CLASSIC_TRANSITIONS[k]['tag'].lower()}, 得到 {ln}"
    check("T2 全部经典切换效果 XML 生成", t2)

    # ---- T3 p14 效果包裹 AlternateContent 且带 fallback ---------------------
    def t3():
        keys = list(P14_TRANSITIONS)
        out = os.path.join(tmp, "t3.pptx")
        per = [{"slide": i + 1, "transition": k, "advance": 1.0}
               for i, k in enumerate(keys)]
        build_pptx([images[0]] * len(keys), out, transition=keys[0],
                   advance=1.0, per_slide=per)
        with zipfile.ZipFile(out) as z:
            for i, k in enumerate(keys):
                x = z.read("ppt/slides/slide%d.xml" % (i + 1)).decode("utf-8")
                assert "mc:AlternateContent" in x, f"{k}: 缺少 AlternateContent"
                assert "mc:Fallback" in x and "<p:fade/>" in x, f"{k}: 缺少 fallback"
                assert P14_TRANSITIONS[k]["tag"] in x, f"{k}: 缺少 p14 子元素"
        # morph 同理；conveyor/morph 的必填属性应被自动补全
        out2 = os.path.join(tmp, "t3b.pptx")
        build_pptx([images[0]], out2, transition="morph", advance=1.0)
        with zipfile.ZipFile(out2) as z:
            x = z.read("ppt/slides/slide1.xml").decode("utf-8")
            assert 'Requires="p159"' in x and "morph" in x, "morph 结构不正确"
            assert 'option="byObject"' in x, "morph 必填 option 未自动补全"
        out3 = os.path.join(tmp, "t3c.pptx")
        build_pptx([images[0]], out3, transition="conveyor", advance=1.0)
        with zipfile.ZipFile(out3) as z:
            x = z.read("ppt/slides/slide1.xml").decode("utf-8")
            assert '<p14:conveyor dir="l"/>' in x, "conveyor 必填 dir 未自动补全"
    check("T3 p14/morph 切换的兼容包裹结构", t3)

    # ---- T4 set-transitions 修改已有文件 ------------------------------------
    def t4():
        import shutil
        out = os.path.join(tmp, "t4.pptx")
        build_pptx(images, out, transition="fade", advance=2.0)
        prs = Presentation(out)
        set_slide_transition(prs.slides[1],
                             {"transition": "push", "direction": "left",
                              "duration": None, "options": None}, 3500)
        set_slide_transition(prs.slides[2], {"transition": "none"}, 4000)
        # 第 4 页沿用已有 advTm
        set_slide_transition(prs.slides[3], {"transition": "wipe", "direction": "up"})
        prs.save(out)
        prs2 = Presentation(out)
        assert _slide_transition_localname(prs2.slides[1]) == "push"
        from pptx.oxml.ns import qn
        tr2 = prs2.slides[1]._element.find(qn("p:transition"))
        assert list(tr2)[0].get("dir") == "l", "push 方向应在子元素上"
        assert tr2.get("advTm") == "3500", "advTm 在 transition 元素上"
        tr3 = prs2.slides[2]._element.find(qn("p:transition"))
        assert len(list(tr3)) == 0 and tr3.get("advTm") == "4000", "none 切换结构不正确"
        tr4 = prs2.slides[3]._element.find(qn("p:transition"))
        assert _localname(list(tr4)[0].tag) == "wipe" and tr4.get("advTm") == "2000", \
            "keep 模式应沿用原 advTm"
    check("T4 set-transitions / none / keep-advance", t4)

    # ---- T5 图片适配几何 -----------------------------------------------------
    def t5():
        def approx(a, b, tol=2):
            return abs(a - b) <= tol

        results = {}
        for mode in ("contain", "cover", "stretch"):
            out = os.path.join(tmp, "fit_%s.pptx" % mode)
            build_pptx([square], out, fit=mode)
            prs = Presentation(out)
            # 音乐未启用，页内只有 1 张图片
            pics = list(prs.slides[0].shapes)
            assert len(pics) == 1, f"{mode}: 应只有 1 个形状, 得到 {len(pics)}"
            pic = pics[0]
            sw, sh_ = prs.slide_width, prs.slide_height
            w, h = pic.width, pic.height
            results[mode] = (pic.left, pic.top, w, h)
            if mode == "stretch":
                assert approx(w, sw) and approx(h, sh_), "stretch 尺寸不正确"
            elif mode == "cover":
                assert w >= sw - 2 and h >= sh_ - 2, "cover 未铺满"
                assert approx(pic.left, (sw - w) // 2) and approx(pic.top, (sh_ - h) // 2), \
                    "cover 未居中"
            else:
                assert w <= sw + 2 and h <= sh_ + 2, "contain 溢出画布"
                fits = (approx(h, sh_) and approx(pic.left, (sw - w) // 2)
                        and approx(pic.top, 0)) or \
                       (approx(w, sw) and approx(pic.top, (sh_ - h) // 2)
                        and approx(pic.left, 0))
                assert fits, "contain 未正确贴边居中"
        # contain 与 cover 尺寸关系：cover 面积 >= contain 面积
        assert results["cover"][2] * results["cover"][3] >= \
               results["contain"][2] * results["contain"][3]
    check("T5 contain/cover/stretch 几何计算", t5)

    # ---- T6 自然排序 ----------------------------------------------------------
    def t6():
        assert natural_key("img10.png") > natural_key("img9.png")
        assert natural_key("IMG2.png") < natural_key("img10.png")
        sub = os.path.join(tmp, "sort")
        os.makedirs(sub, exist_ok=True)
        for nm in ("b.png", "a10.png", "a2.png", "a1.png", "note.txt"):
            with open(os.path.join(sub, nm), "wb") as f:
                f.write(make_png_bytes(2, 2, (1, 2, 3)) if nm.endswith(".png") else b"x")
        got = collect_images([sub])
        names = [os.path.basename(p) for p in got]
        assert names == ["a1.png", "a2.png", "a10.png", "b.png"], f"排序错误: {names}"
    check("T6 自然排序与目录收集", t6)

    # ---- T7 逐页 JSON 覆盖 -----------------------------------------------------
    def t7():
        tjson = os.path.join(tmp, "tr.json")
        with open(tjson, "w", encoding="utf-8") as f:
            json.dump([
                {"transition": "wipe", "direction": "up", "advance": 1.5},
                {"slide": 3, "transition": "zoom", "direction": "in", "advance": 2.5},
            ], f)
        out = os.path.join(tmp, "t7.pptx")
        build_pptx(images, out, transition="fade", advance=9.0,
                   per_slide=parse_transitions_json(tjson))
        prs = Presentation(out)
        assert _slide_transition_localname(prs.slides[0]) == "wipe"
        assert _slide_advTm(prs.slides[0]) == "1500"
        assert _slide_transition_localname(prs.slides[1]) == "fade", "未指定的页应保持默认"
        assert _slide_advTm(prs.slides[1]) == "9000"
        assert _slide_transition_localname(prs.slides[2]) == "zoom"
        assert _slide_advTm(prs.slides[2]) == "2500"
        from pptx.oxml.ns import qn
        tr1 = prs.slides[0]._element.find(qn("p:transition"))
        assert list(tr1)[0].get("dir") == "u", "wipe 方向应在子元素上"
    check("T7 transitions-json 逐页覆盖", t7)

    # ---- T8 错误处理 ------------------------------------------------------------
    def t8():
        for bad in ("no_such_effect", "fade2"):
            try:
                lookup_transition(bad)
                raise AssertionError(f"应拒绝未知效果 {bad}")
            except SkillError:
                pass
        try:
            build_pptx([], os.path.join(tmp, "never.pptx"))
            raise AssertionError("空图片列表应报错")
        except SkillError:
            pass
        try:
            transition_attrs("classic", "fade", CLASSIC_TRANSITIONS["fade"],
                             "left", None)
            raise AssertionError("fade 不应接受方向 left")
        except SkillError:
            pass
        try:
            transition_attrs("classic", "wheel", CLASSIC_TRANSITIONS["wheel"],
                             None, {"spokes": "5"})
            raise AssertionError("spokes=5 应被拒绝")
        except SkillError:
            pass
    check("T8 参数校验与错误处理", t8)

    # ---- T9 CLI 冒烟（argparse 接线） -------------------------------------------
    def t9():
        out = os.path.join(tmp, "t9.pptx")
        rc = main(["build", "--images", images[0], "--output", out,
                   "--transition", "dissolve", "--advance", "3"])
        assert rc == 0, f"CLI build 返回 {rc}"
        assert os.path.isfile(out)
        prs = Presentation(out)
        assert _slide_transition_localname(prs.slides[0]) == "dissolve"
        assert _slide_advTm(prs.slides[0]) == "3000"
    check("T9 CLI 冒烟测试", t9)

    # ---- T10 音频循环开关与音量 ---------------------------------------------------
    def t10():
        out = os.path.join(tmp, "t10.pptx")
        build_pptx(images, out, music=wav, transition="fade", advance=1.0,
                   music_loop=False, music_volume=50)
        with zipfile.ZipFile(out) as z:
            s1 = z.read("ppt/slides/slide1.xml").decode("utf-8")
            assert 'vol="50000"' in s1, "音量未正确写入"
            assert 'repeatCount="indefinite"' not in s1, "不应循环"
        out2 = os.path.join(tmp, "t10b.pptx")
        build_pptx(images, out2, music=wav, transition="fade", advance=1.0)
        with zipfile.ZipFile(out2) as z:
            s1 = z.read("ppt/slides/slide1.xml").decode("utf-8")
            assert 'repeatCount="indefinite"' in s1, "默认应循环"
            assert 'numSld="999"' in s1, "默认应跨全部页"
    check("T10 音乐循环/音量/跨页参数", t10)

    # ---- T11 p14:dur 精确时长结构 -------------------------------------------------
    def t11():
        out = os.path.join(tmp, "t11.pptx")
        build_pptx([images[0]], out, transition="fade", duration=0.7, advance=1.0)
        with zipfile.ZipFile(out) as z:
            x = z.read("ppt/slides/slide1.xml").decode("utf-8")
            assert 'p14:dur="700"' in x, "缺少精确切换时长"
            assert 'spd="fast"' in x, "spd 降级档缺失"
        out2 = os.path.join(tmp, "t11b.pptx")
        build_pptx([images[0]], out2, transition="fade", duration=2.0, advance=1.0)
        with zipfile.ZipFile(out2) as z:
            x = z.read("ppt/slides/slide1.xml").decode("utf-8")
            assert 'p14:dur="2000"' in x and 'spd="slow"' in x
    check("T11 切换时长（p14:dur + spd 降级）", t11)

    return failures


def cmd_selftest(args):
    tmp = tempfile.mkdtemp(prefix="pptx_builder_selftest_")
    log(f"[信息] 自检目录: {tmp}")
    log("[信息] 开始离线自检（不依赖 PowerPoint）...")
    failures = _run_tests(tmp)
    total = 11
    passed = total - len(failures)
    summary = {"passed": passed, "failed": len(failures), "total": total,
               "workdir": tmp,
               "failures": [{"name": n, "error": repr(e)} for n, e in failures]}
    if failures:
        log(f"[结果] {passed}/{total} 通过，存在失败项！")
        if args.json:
            json_out(summary)
        return 1
    log(f"[结果] 全部 {total} 项自检通过。")
    if not args.keep:
        try:
            import shutil
            shutil.rmtree(tmp, ignore_errors=True)
            summary["workdir"] = None
        except Exception:
            pass
    if args.json:
        json_out(summary)
    return 0


# --------------------------------------------------------------------------
# argparse 接线
# --------------------------------------------------------------------------

def _add_transition_args(p, with_default_advance=True):
    p.add_argument("--transition", default="fade" if with_default_advance else None,
                   help="切换效果名（默认 fade；见 list-transitions）")
    p.add_argument("--direction", default=None,
                   help="方向，如 left/right/up/down、horz/vert、in/out、throughblack")
    p.add_argument("--transition-duration", type=float, default=None, metavar="秒",
                   help="切换动画时长（秒），精确时长需 PowerPoint 2010+")
    if with_default_advance:
        p.add_argument("--advance", type=float, default=5.0, metavar="秒",
                       help="每页停留秒数（自动换片；导出视频节奏。默认 5.0，0=不自动换片）")
    else:
        p.add_argument("--advance", type=float, default=None, metavar="秒",
                       help="每页停留秒数（缺省=沿用文件中已有值）")
    p.add_argument("--opt", action="append", default=None, metavar="键=值",
                   help="切换附加参数，可重复。如 --opt spokes=3 --opt pattern=hexagon")
    p.add_argument("--transitions-json", default=None, metavar="文件",
                   help="逐页切换 JSON 文件（格式见脚本头部文档）")


def _add_media_args(p):
    p.add_argument("--music", default=None, metavar="文件",
                   help="背景音乐文件（mp3/m4a/wav/wma/aac/flac）")
    p.add_argument("--music-no-loop", action="store_true",
                   help="背景音乐不循环（默认循环到结束）")
    p.add_argument("--music-volume", type=int, default=100, metavar="1-100",
                   help="背景音乐音量（默认 100）")
    p.add_argument("--auto-fit-music", action="store_true",
                   help="按音乐时长自动分配每页停留时间（需 mutagen；覆盖 --advance）")


def _add_image_args(p):
    p.add_argument("--images", nargs="+", metavar="路径",
                   help="图片：文件/目录/glob 通配（png jpg jpeg gif bmp tif）")
    p.add_argument("--recursive", action="store_true", help="目录递归收集")
    p.add_argument("--fit", choices=["contain", "cover", "stretch"], default="contain",
                   help="图片适配：contain 等比含边距(默认) / cover 等比铺满裁边 / stretch 拉伸")
    p.add_argument("--background", default="black", metavar="颜色",
                   help="背景色（颜色名或 #RRGGBB，默认 black）")
    p.add_argument("--slide-size", default="16:9", metavar="尺寸",
                   help="幻灯片尺寸：16:9(默认)/4:3/16:10/1280x720")


def _build_parser():
    p = argparse.ArgumentParser(
        prog="pptx_video_builder.py",
        description="批量图片 + 背景音乐 -> PPTX -> MP4 视频（详见脚本头部文档）",
        formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="command", metavar="命令")

    pb = sub.add_parser("build", help="图片(+音乐) -> PPTX")
    _add_image_args(pb)
    _add_media_args(pb)
    _add_transition_args(pb)
    pb.add_argument("--output", required=True, metavar="文件", help="输出 .pptx 路径")
    pb.add_argument("--json", action="store_true", help="以 JSON 输出结果(stdout)")
    pb.set_defaults(func=cmd_build)

    ps = sub.add_parser("set-transitions", help="修改已有 PPTX 的切换方式")
    ps.add_argument("--pptx", required=True, metavar="文件", help="目标 .pptx")
    _add_transition_args(ps, with_default_advance=False)
    ps.add_argument("--json", action="store_true")
    ps.set_defaults(func=cmd_set_transitions)

    pl = sub.add_parser("list-transitions", help="列出切换效果并实测本机支持情况")
    pl.add_argument("--no-verify", action="store_true",
                    help="跳过本机 PowerPoint 实测（仅输出理论目录）")
    pl.add_argument("--json", action="store_true")
    pl.set_defaults(func=cmd_list_transitions)

    pe = sub.add_parser("export", help="PPTX -> MP4（需要本机 PowerPoint 2010+）")
    pe.add_argument("--pptx", required=True, metavar="文件")
    pe.add_argument("--output", required=True, metavar="文件", help="输出 .mp4 路径")
    pe.add_argument("--resolution", type=int, default=720, choices=[480, 720, 1080],
                    help="视频分辨率（默认 720）")
    pe.add_argument("--fps", type=int, default=30, choices=[24, 30, 60])
    pe.add_argument("--quality", type=int, default=85, metavar="1-100",
                    help="视频质量（默认 85）")
    pe.add_argument("--timeout", type=int, default=1800, metavar="秒",
                    help="导出超时（默认 1800）")
    pe.add_argument("--slide-seconds", type=int, default=None, metavar="秒",
                    help="无换片时间页面的默认时长（默认沿用 -1=PowerPoint 默认）")
    pe.add_argument("--json", action="store_true")
    pe.set_defaults(func=cmd_export)

    pa = sub.add_parser("all", help="build + export 一条龙")
    _add_image_args(pa)
    _add_media_args(pa)
    _add_transition_args(pa)
    pa.add_argument("--output", required=True, metavar="文件", help="输出 .mp4 路径")
    pa.add_argument("--pptx", default=None, metavar="文件",
                    help="中间 PPTX 路径（默认与视频同名的 .pptx）")
    pa.add_argument("--resolution", type=int, default=720, choices=[480, 720, 1080])
    pa.add_argument("--fps", type=int, default=30, choices=[24, 30, 60])
    pa.add_argument("--quality", type=int, default=85, metavar="1-100")
    pa.add_argument("--timeout", type=int, default=1800, metavar="秒")
    pa.add_argument("--slide-seconds", type=int, default=None)
    pa.add_argument("--json", action="store_true")
    pa.set_defaults(func=cmd_all)

    pt = sub.add_parser("selftest", help="离线自检（不需要 PowerPoint）")
    pt.add_argument("--keep", action="store_true", help="保留自检临时目录")
    pt.add_argument("--json", action="store_true")
    pt.set_defaults(func=cmd_selftest)
    return p


def main(argv=None):
    parser = _build_parser()
    args = parser.parse_args(argv)
    if not getattr(args, "command", None):
        parser.print_help()
        return 2
    try:
        return args.func(args)
    except SkillError as e:
        log(f"[错误] {e}")
        return e.exit_code
    except KeyboardInterrupt:
        log("[中断] 用户取消")
        return 130


if __name__ == "__main__":
    sys.exit(main())
